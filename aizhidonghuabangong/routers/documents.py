# routers/documents.py — 文档生成接口
# 读取客户+订单数据，生成固定模板 Word 文档，支持下载

import io
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse

from sqlalchemy import select

from database import get_db
from models import Customer, Order
from auth import get_current_user

router = APIRouter(prefix="/api/documents", tags=["文档生成"])


@router.get("/generate/{order_id}")
def generate_document(
    order_id: int,
    db = Depends(get_db),
    _user=Depends(get_current_user)
):
    """
    根据订单ID生成 Word 文档
    文档包含客户信息和订单详情，使用固定模板
    """
    # 查询订单
    order_result = db.execute(select(Order).where(Order.id == order_id))
    order = order_result.scalar_one_or_none()
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")

    # 查询客户
    cust_result = db.execute(select(Customer).where(Customer.id == order.customer_id))
    customer = cust_result.scalar_one_or_none()
    if not customer:
        raise HTTPException(status_code=404, detail="关联客户不存在")

    # 使用 python-docx 生成 Word 文档
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)

    # ===== 文档标题 =====
    title = doc.add_heading('订单确认单', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ===== 订单基本信息 =====
    doc.add_heading('订单信息', level=2)
    table = doc.add_table(rows=5, cols=2, style='Table Grid')

    # 填充订单信息
    cells = [
        ('订单编号', order.order_no),
        ('订单金额', f'¥{order.amount:,.2f}'),
        ('商品信息', order.product_info or '无'),
        ('交货时间', str(order.delivery_date) if order.delivery_date else '未设定'),
        ('订单状态', order.status),
    ]
    for i, (label, value) in enumerate(cells):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = str(value)

    # ===== 客户信息 =====
    doc.add_heading('客户信息', level=2)
    table2 = doc.add_table(rows=5, cols=2, style='Table Grid')

    cells2 = [
        ('客户名称', customer.name),
        ('联系人', customer.contact or '无'),
        ('联系电话', customer.phone or '无'),
        ('地址', customer.address or '无'),
        ('备注', customer.notes or '无'),
    ]
    for i, (label, value) in enumerate(cells2):
        table2.rows[i].cells[0].text = label
        table2.rows[i].cells[1].text = str(value)

    # ===== 页脚说明 =====
    doc.add_paragraph('')
    doc.add_paragraph('此文档由系统自动生成，盖章签字后生效。')

    # 保存到内存
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    # 返回文件流
    filename = f"order_{order.order_no}.docx"
    encoded_filename = quote(filename)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f"attachment; filename={filename}; filename*=UTF-8''{encoded_filename}"
        }
    )
